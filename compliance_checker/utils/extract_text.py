# Text extraction utilities for parsing PDF, DOCX, and plain text documents.

import os
import io
import PyPDF2
import docx


def extract_text(uploaded_file) -> str:
    """
    Extracts all textual content from an uploaded file (.pdf, .docx, or .txt).

    Parameters:
        uploaded_file: A Streamlit UploadedFile object, file-like object, or file path.

    Returns:
        str: The extracted full text as a single string, or an empty string if an error occurs.
    """
    # Defensive check: ensure a valid file was provided
    if uploaded_file is None:
        return ""

    # Determine filename to check the extension
    # Streamlit UploadedFile objects provide the uploaded filename via the '.name' attribute.
    if hasattr(uploaded_file, "name"):
        filename = uploaded_file.name
    else:
        filename = str(uploaded_file)

    # Extract the file extension in lowercase (e.g., '.pdf', '.docx', '.txt')
    _, extension = os.path.splitext(filename)
    extension = extension.lower()

    # Reset file pointer to the beginning in case the file was read previously
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    try:
        # Case 1: PDF Document
        if extension == ".pdf":
            # --- How PyPDF2 Works ---
            # 1. PyPDF2 reads binary PDF data from a file stream.
            # 2. PyPDF2.PdfReader parses the internal PDF tree structure (catalog, pages, objects).
            # 3. We can inspect 'reader.pages', which is a list-like collection of page objects.
            # 4. Calling 'page.extract_text()' extracts the text characters and layout from each individual page.
            reader = PyPDF2.PdfReader(uploaded_file)
            page_texts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    page_texts.append(text)

            # Join all page texts together separated by newline characters
            extracted_text = "\n".join(page_texts)

        # Case 2: Word Document (.docx)
        elif extension == ".docx":
            # --- How python-docx Works ---
            # 1. Word (.docx) files are zipped XML archives containing styles, settings, and body text.
            # 2. docx.Document() unpacks this document structure from a file path or file-like binary stream.
            # 3. 'doc.paragraphs' gives access to all paragraph elements in reading order.
            # 4. Each 'paragraph.text' retrieves the raw string inside that paragraph run.
            doc = docx.Document(uploaded_file)
            paragraph_texts = [paragraph.text for paragraph in doc.paragraphs]

            # Join all paragraph texts together separated by newline characters
            extracted_text = "\n".join(paragraph_texts)

        # Case 3: Plain Text File (.txt)
        elif extension == ".txt":
            # For plain text, we read the raw byte content and decode it into a Python string.
            if hasattr(uploaded_file, "read"):
                raw_bytes = uploaded_file.read()
                # Most text files use UTF-8 encoding. If UTF-8 fails, we fall back to latin-1.
                try:
                    extracted_text = raw_bytes.decode("utf-8")
                except UnicodeDecodeError:
                    extracted_text = raw_bytes.decode("latin-1", errors="replace")
            else:
                # If uploaded_file is already a string path
                with open(filename, "r", encoding="utf-8", errors="replace") as f:
                    extracted_text = f.read()

        # Unsupported file extension
        else:
            print(f"Error: Unsupported file extension '{extension}'. Supported formats: .pdf, .docx, .txt")
            return ""

        # Reset stream position again after reading so subsequent steps can reuse it
        if hasattr(uploaded_file, "seek"):
            uploaded_file.seek(0)

        return extracted_text

    except Exception as error:
        # Gracefully handle unexpected errors (corrupted files, malformed archives, etc.)
        print(f"Error extracting text from '{filename}': {error}")
        return ""

